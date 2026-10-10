"""存库第二步：PDF/DOCX/文本解析，再按标题与页码切片；保留引用位置。"""
import io
import re
from dataclasses import dataclass


@dataclass
class ParsedChunk:
    content: str
    page: int | None
    heading: str


def step1_parse_document(filename: str, data: bytes) -> list[tuple[str, int | None]]:
    """返回文本块和页码；扫描件未配置 OCR 时明确失败，不发布零切片。"""
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ValueError("加密 PDF 需要先解除密码")
            pages = [(page.get_text(), i + 1) for i, page in enumerate(document)]
        if not any(text.strip() for text, _ in pages):
            raise ValueError("扫描 PDF 没有文字层；当前未配置 OCR，请上传带文字层的 PDF")
        return pages
    if suffix == "docx":
        from docx import Document
        document = Document(io.BytesIO(data))
        texts = [p.text for p in document.paragraphs]
        texts.extend(" | ".join(cell.text for cell in row.cells) for t in document.tables for row in t.rows)
        return [("\n".join(texts), None)]
    return [(data.decode("utf-8-sig"), None)]


def step2_split_into_citable_chunks(pages: list[tuple[str, int | None]], size=600, overlap=80):
    """按页保留来源；段内切片有重叠，短文本也能形成至少一片。"""
    if size <= overlap or overlap < 0:
        raise ValueError("切片长度必须大于重叠长度")
    chunks = []
    for text, page in pages:
        text = re.sub(r"[ \t]+", " ", text).strip()
        heading = ""
        for paragraph in re.split(r"\n(?=#{1,6} )", text):
            if paragraph.startswith("#"):
                heading = paragraph.splitlines()[0].lstrip("# ")[:200]
            for start in range(0, len(paragraph), size - overlap):
                content = paragraph[start:start + size].strip()
                if content:
                    chunks.append(ParsedChunk(content, page, heading))
                if start + size >= len(paragraph):
                    break
    if not chunks:
        raise ValueError("文档没有可检索的有效文本")
    return chunks
